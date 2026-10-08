"""Rendu de la revue générée (Markdown) en PDF (WeasyPrint) et DOCX (python-docx).

Le document exporté est celui que le client diffuse à sa direction : en-tête de
publication (titre, organisation, date), la revue elle-même, puis une annexe « Points clés
détectés » regroupée par catégorie. Le rendu est mis en cache sur le modèle : un même
format n'est généré qu'une fois.

Sécurité : le texte de la revue est rédigé par l'IA à partir d'articles fournis par le
client — il peut donc contenir du HTML hostile recopié d'une coupure. Le Markdown est rendu
sans HTML brut (`<` échappé) et WeasyPrint n'a le droit de charger AUCUNE ressource
(`_deny_all_fetcher`) : sans ça, un `<link rel="attachment" href="file:///…">` glissé
dans la revue embarquerait un fichier du serveur dans le PDF.
"""
import html
import io
import logging
import re
from collections import OrderedDict

import markdown as md
from django.core.files.base import ContentFile
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

logger = logging.getLogger(__name__)

_MONTHS_FR = (
    "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre", "décembre",
)

PDF_STYLE = """
@page {
  size: A4;
  margin: 20mm 18mm 22mm 18mm;
  @bottom-left { content: "Revue de Presse · %(org)s"; font-size: 7.5pt; color: #94a3b8; font-family: %(sans)s; }
  @bottom-right { content: "Page " counter(page) " / " counter(pages); font-size: 7.5pt; color: #94a3b8; font-family: %(sans)s; }
}
body { font-family: %(sans)s; color: #1e293b; font-size: 10.5pt; line-height: 1.6; }
.masthead { border-bottom: 3px solid #0f172a; padding-bottom: 10pt; margin-bottom: 22pt; }
.masthead .kicker { font-size: 8pt; letter-spacing: 2.5pt; text-transform: uppercase; color: #4f46e5; font-weight: bold; }
.masthead .title { font-family: %(serif)s; font-size: 24pt; line-height: 1.15; margin: 6pt 0 6pt 0; color: #0f172a; }
.masthead .meta { font-size: 9pt; color: #64748b; }
.masthead .bar { height: 5pt; width: 60pt; background: #fbbf24; margin-top: 10pt; }
h1, h2, h3 { font-family: %(serif)s; color: #0f172a; line-height: 1.25; page-break-after: avoid; }
.review h1 { font-size: 17pt; margin: 18pt 0 8pt; }
.review h2 { font-size: 14pt; margin: 18pt 0 6pt; padding-bottom: 3pt; border-bottom: 1px solid #e2e8f0; }
.review h3 { font-size: 11.5pt; margin: 12pt 0 4pt; }
.review p { margin: 0 0 8pt; text-align: justify; }
.review ul, .review ol { margin: 0 0 8pt; padding-left: 16pt; }
.review li { margin-bottom: 3pt; }
.review strong { color: #0f172a; }
.review blockquote { margin: 10pt 0; padding: 6pt 10pt; border-left: 3pt solid #fbbf24; background: #fffbeb; font-style: italic; }
.review table { border-collapse: collapse; width: 100%%; margin: 8pt 0; font-size: 9.5pt; }
.review th, .review td { border-bottom: 1px solid #e2e8f0; padding: 4pt 6pt; text-align: left; vertical-align: top; }
.annex { page-break-before: always; }
.annex > h2 { font-size: 16pt; margin: 0 0 4pt; }
.annex .intro { color: #64748b; font-size: 9pt; margin-bottom: 14pt; }
.annex h3 { font-size: 12pt; margin: 16pt 0 8pt; }
.annex h3 .dot { display: inline-block; width: 8pt; height: 8pt; border-radius: 4pt; margin-right: 5pt; }
.annex h3 .count { color: #94a3b8; font-weight: normal; font-size: 10pt; }
.hl { border-left: 3pt solid #cbd5e1; padding: 2pt 0 2pt 10pt; margin: 0 0 10pt; page-break-inside: avoid; }
.hl .excerpt { font-family: %(serif)s; font-style: italic; margin: 0 0 3pt; color: #0f172a; }
.hl .why { margin: 0 0 2pt; color: #475569; font-size: 9.5pt; }
.hl .src { margin: 0; color: #94a3b8; font-size: 8pt; }
.disclaimer { margin-top: 18pt; padding-top: 8pt; border-top: 1px solid #e2e8f0; font-size: 8pt; color: #94a3b8; }
"""
SERIF_STACK = '"P052", "Palatino Linotype", Palatino, Georgia, "DejaVu Serif", serif'
SANS_STACK = '"Nimbus Sans", "Helvetica Neue", Helvetica, Arial, "DejaVu Sans", sans-serif'

DISCLAIMER = (
    "Synthèse et points clés générés automatiquement par l'IA à partir des documents fournis. "
    "Ils ne constituent pas un avis juridique : vérifiez les passages sensibles avant toute décision."
)


def get_or_render_export(review, *, fmt):
    if fmt == "pdf":
        if not review.pdf_file:
            _render_pdf(review)
        return review.pdf_file
    if fmt == "docx":
        if not review.docx_file:
            _render_docx(review)
        return review.docx_file
    raise ValueError(f"Format d'export non supporté : {fmt}")


# --------------------------------------------------------------------------- données


def _french_date(moment):
    moment = timezone.localtime(moment or timezone.now())
    return f"{moment.day} {_MONTHS_FR[moment.month - 1]} {moment.year}"


def _context(review):
    job = review.job
    return {
        "title": job.title or "Revue de presse",
        "organization": job.organization.name,
        "date": _french_date(review.generated_at or job.created_at),
        "documents": job.documents.count(),
    }


def _grouped_highlights(review):
    """Points clés regroupés par catégorie, dans l'ordre d'affichage des catégories."""
    from apps.press_review.models import Highlight

    groups = OrderedDict()
    highlights = (
        Highlight.objects.filter(job=review.job)
        .select_related("category", "source_document")
        .order_by("category__sort_order", "category__name", "order")
    )
    for highlight in highlights:
        groups.setdefault(highlight.category, []).append(highlight)
    return groups


def _safe_color(value, default="#64748B"):
    return value if re.fullmatch(r"#[0-9a-fA-F]{6}", value or "") else default


def _source_label(highlight):
    if highlight.source_document is None:
        return ""
    label = highlight.source_document.original_filename
    if highlight.page_number and highlight.source_document.kind == "pdf":
        label += f" · page {highlight.page_number}"
    return label


# --------------------------------------------------------------------------- PDF


def _deny_all_fetcher(url, *args, **kwargs):
    raise ValueError(f"Ressource externe refusée dans l'export : {url[:80]}")


def _markdown_to_safe_html(text):
    # `<` échappé : aucune balise HTML brute ne peut passer du texte de la revue au PDF.
    return md.markdown((text or "").replace("<", "&lt;"), extensions=["extra", "sane_lists"])


def _build_pdf_html(review):
    ctx = _context(review)
    esc = html.escape
    parts = [
        f"<header class='masthead'><div class='kicker'>Revue de presse</div>"
        f"<h1 class='title'>{esc(ctx['title'])}</h1>"
        f"<div class='meta'>{esc(ctx['organization'])} · {esc(ctx['date'])} · "
        f"{ctx['documents']} document{'s' if ctx['documents'] > 1 else ''} analysé{'s' if ctx['documents'] > 1 else ''}</div>"
        "<div class='bar'></div></header>",
        f"<section class='review'>{_markdown_to_safe_html(review.summary_markdown)}</section>",
    ]

    groups = _grouped_highlights(review)
    if groups:
        parts.append(
            "<section class='annex'><h2>Points clés détectés</h2>"
            "<p class='intro'>Passages repérés par l'IA dans les documents, classés selon vos catégories de détection.</p>"
        )
        for category, items in groups.items():
            color = _safe_color(category.color)
            parts.append(
                f"<h3><span class='dot' style='background:{color}'></span>{esc(category.name)} "
                f"<span class='count'>({len(items)})</span></h3>"
            )
            for highlight in items:
                source = _source_label(highlight)
                parts.append(
                    f"<div class='hl' style='border-left-color:{color}'>"
                    f"<p class='excerpt'>« {esc(highlight.excerpt)} »</p>"
                    + (f"<p class='why'>{esc(highlight.explanation)}</p>" if highlight.explanation else "")
                    + (f"<p class='src'>{esc(source)}</p>" if source else "")
                    + "</div>"
                )
        parts.append("</section>")

    parts.append(f"<p class='disclaimer'>{esc(DISCLAIMER)}</p>")
    # Nom d'organisation injecté dans une chaîne CSS (pied de page) : ni guillemet ni antislash.
    footer_org = re.sub(r'["\\\n\r]', "", ctx["organization"])
    style = PDF_STYLE % {"org": footer_org, "serif": SERIF_STACK, "sans": SANS_STACK}
    return f"<html lang='fr'><head><meta charset='utf-8'><style>{style}</style></head><body>{''.join(parts)}</body></html>"


def _render_pdf(review):
    # Import paresseux : WeasyPrint a besoin de libs système natives (Pango/cairo,
    # présentes dans le Dockerfile) absentes de certains postes de dev (Windows sans
    # GTK). Un import en haut de fichier casserait aussi l'export DOCX, qui n'en a pas
    # besoin, dès que ces libs manquent.
    from weasyprint import HTML

    pdf_bytes = HTML(string=_build_pdf_html(review), url_fetcher=_deny_all_fetcher).write_pdf()
    review.pdf_file.save(f"revue-{review.job_id}.pdf", ContentFile(pdf_bytes), save=False)
    review.save(update_fields=["pdf_file", "updated_at"])


# --------------------------------------------------------------------------- DOCX

_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_)")
_ORDERED_ITEM = re.compile(r"^\d+[.)]\s+")


def _add_inline_runs(paragraph, text):
    """**gras** et *italique* Markdown -> vraies mises en forme Word (au lieu d'astérisques)."""
    for token in _INLINE.split(text):
        if not token:
            continue
        if (token.startswith("**") and token.endswith("**")) or (token.startswith("__") and token.endswith("__")):
            paragraph.add_run(token[2:-2]).bold = True
        elif (token.startswith("*") and token.endswith("*")) or (token.startswith("_") and token.endswith("_")):
            paragraph.add_run(token[1:-1]).italic = True
        else:
            paragraph.add_run(token)


def _hex_to_rgb(value):
    value = _safe_color(value).lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _render_docx(review):
    ctx = _context(review)
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)

    kicker = document.add_paragraph()
    kicker_run = kicker.add_run("REVUE DE PRESSE")
    kicker_run.bold = True
    kicker_run.font.size = Pt(9)
    kicker_run.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
    document.add_heading(ctx["title"], level=0)
    meta = document.add_paragraph(f"{ctx['organization']} · {ctx['date']}")
    meta.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    for line in (review.summary_markdown or "").splitlines():
        stripped = line.strip()
        if not stripped or stripped in ("---", "***"):
            continue
        if stripped.startswith("### "):
            document.add_heading(stripped[4:].strip("*"), level=3)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:].strip("*"), level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:].strip("*"), level=1)
        elif stripped.startswith(("- ", "* ", "+ ")):
            _add_inline_runs(document.add_paragraph(style="List Bullet"), stripped[2:])
        elif _ORDERED_ITEM.match(stripped):
            _add_inline_runs(document.add_paragraph(style="List Number"), _ORDERED_ITEM.sub("", stripped))
        elif stripped.startswith(">"):
            quote = document.add_paragraph(style="Intense Quote")
            _add_inline_runs(quote, stripped.lstrip("> "))
        else:
            _add_inline_runs(document.add_paragraph(), stripped)

    groups = _grouped_highlights(review)
    if groups:
        document.add_page_break()
        document.add_heading("Points clés détectés", level=1)
        intro = document.add_paragraph(
            "Passages repérés par l'IA dans les documents, classés selon vos catégories de détection."
        )
        intro.runs[0].italic = True
        for category, items in groups.items():
            heading = document.add_heading(level=2)
            name_run = heading.add_run(f"{category.name} ({len(items)})")
            name_run.font.color.rgb = _hex_to_rgb(category.color)
            for highlight in items:
                excerpt = document.add_paragraph(style="List Bullet")
                excerpt_run = excerpt.add_run(f"« {highlight.excerpt} »")
                excerpt_run.italic = True
                if highlight.explanation:
                    explanation = document.add_paragraph(highlight.explanation)
                    explanation.paragraph_format.left_indent = Pt(18)
                    explanation.runs[0].font.size = Pt(10)
                source = _source_label(highlight)
                if source:
                    src = document.add_paragraph(source)
                    src.paragraph_format.left_indent = Pt(18)
                    src.runs[0].font.size = Pt(8.5)
                    src.runs[0].font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    disclaimer = document.add_paragraph(DISCLAIMER)
    disclaimer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    disclaimer.runs[0].font.size = Pt(8.5)
    disclaimer.runs[0].font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    buffer = io.BytesIO()
    document.save(buffer)
    review.docx_file.save(f"revue-{review.job_id}.docx", ContentFile(buffer.getvalue()), save=False)
    review.save(update_fields=["docx_file", "updated_at"])
