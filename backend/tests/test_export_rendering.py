"""Contenu des exports PDF/DOCX : mise en page, annexe des points clés, et durcissement
contre le HTML hostile recopié depuis un article (la revue est rédigée par l'IA)."""
import io

import pytest
from django.utils import timezone
from docx import Document

from apps.ai_pipeline.export import _build_pdf_html, _deny_all_fetcher, _render_docx
from apps.press_review.models import GeneratedReview, Highlight, PressReviewJob, SourceDocument
from tests.factories import DetectionCategoryFactory, OrganizationFactory

REVIEW_MD = (
    "# Synthèse du jour\n\n"
    "Le **gouvernement** annonce une *réforme* fiscale.\n\n"
    "## Points marquants\n\n"
    "- Hausse de la TVA\n"
    "1. Calendrier serré\n"
    "> Citation d'un ministre\n\n"
    "Piège : <link rel=\"attachment\" href=\"file:///etc/passwd\"><script>alert(1)</script>\n"
)


@pytest.fixture
def ready_review(db):
    org = OrganizationFactory(name="Groupe Médias Douala")
    job = PressReviewJob.objects.create(organization=org, title="Revue du lundi", status="REVIEW_READY")
    document = SourceDocument.objects.create(
        organization=org, job=job, file="organizations/x/source_documents/a.pdf", original_filename="cameroon-tribune.pdf"
    )
    legal = DetectionCategoryFactory(name="Risque juridique", color="#DC2626", sort_order=1)
    intox = DetectionCategoryFactory(name="Intox potentielle", color="#D97706", sort_order=2)
    Highlight.objects.create(
        job=job, category=legal, source_document=document, excerpt="Le directeur est accusé sans preuve",
        explanation="Diffamation possible <b>grave</b>", page_number=2, order=0,
    )
    Highlight.objects.create(job=job, category=intox, excerpt="Les sanctions atteindraient des milliards", order=1)
    return GeneratedReview.objects.create(job=job, summary_markdown=REVIEW_MD, generated_at=timezone.now())


@pytest.mark.django_db
def test_pdf_html_has_masthead_review_and_highlights_annex(ready_review):
    html = _build_pdf_html(ready_review)

    assert "Revue du lundi" in html
    assert "Groupe Médias Douala" in html
    assert "<strong>gouvernement</strong>" in html
    assert "Points clés détectés" in html
    assert "Risque juridique" in html and "Intox potentielle" in html
    assert html.index("Risque juridique") < html.index("Intox potentielle")  # ordre des catégories
    assert "cameroon-tribune.pdf · page 2" in html


@pytest.mark.django_db
def test_pdf_html_never_contains_raw_html_from_the_review_or_highlights(ready_review):
    html = _build_pdf_html(ready_review)

    assert "<link rel=" not in html
    assert "<script>" not in html
    assert "&lt;link" in html  # affiché comme du texte, pas interprété
    assert "<b>grave</b>" not in html


def test_pdf_renderer_refuses_every_external_resource():
    with pytest.raises(ValueError):
        _deny_all_fetcher("file:///etc/passwd")
    with pytest.raises(ValueError):
        _deny_all_fetcher("https://exemple.com/image.png")


@pytest.mark.django_db
def test_docx_turns_markdown_emphasis_into_real_formatting_and_adds_the_annex(ready_review, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    _render_docx(ready_review)

    ready_review.refresh_from_db()
    document = Document(io.BytesIO(ready_review.docx_file.read()))
    texts = [p.text for p in document.paragraphs]

    assert "Revue du lundi" in texts
    assert not any("**" in text for text in texts)
    bold_runs = [run.text for p in document.paragraphs for run in p.runs if run.bold]
    assert "gouvernement" in bold_runs
    assert "Points clés détectés" in texts
    assert "« Le directeur est accusé sans preuve »" in texts
