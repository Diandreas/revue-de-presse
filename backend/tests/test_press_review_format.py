"""Format « revue de presse » par défaut (modèle fourni par le client) : bandeau daté,
rubriques en couleur, une notice par journal qui s'ouvre sur le titre de sa une."""
import io
from datetime import datetime
from zoneinfo import ZoneInfo

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document
from rest_framework_simplejwt.tokens import RefreshToken

from apps.ai_pipeline import tasks as pipeline_tasks
from apps.ai_pipeline.export import _build_pdf_html, _render_docx
from apps.ai_pipeline.mistral_client import FakeMistralClient
from apps.ai_pipeline.press_format import (
    build_structured,
    entry_paragraph,
    long_french_date,
    match_rubrique,
    present,
    to_markdown,
)
from apps.ai_pipeline.tasks import run_press_review_pipeline
from apps.press_review.models import DetectionCategory, GeneratedReview, PipelineStepRun, PressReviewJob, SourceDocument
from tests.pdf_fixtures import make_minimal_pdf_bytes

DOUALA = ZoneInfo("Africa/Douala")


@pytest.mark.parametrize(
    "written, expected",
    [
        ("generale", "generale"),
        ("LA PRESSE GÉNÉRALE", "generale"),
        ("Presse économique et financière", "economie"),
        ("economie", "economie"),
        ("PRESSE RÉGIONALE, SÉCURITÉ & DIPLOMATIE", "regionale"),
        ("santé", "societe"),
        ("Politique", "politique"),
        ("Décès et hommages", "deces"),
        ("sport", "sport"),
        ("", "generale"),
        ("n'importe quoi", "generale"),
    ],
)
def test_rubrique_is_recognised_however_the_model_writes_it(written, expected):
    assert match_rubrique(written) == expected


def test_masthead_date_is_written_in_french():
    assert long_french_date(datetime(2026, 3, 5, 9, tzinfo=DOUALA)) == "Jeudi 5 mars 2026"
    assert long_french_date(datetime(2026, 10, 1, 9, tzinfo=DOUALA)) == "Jeudi 1er octobre 2026"


def test_structure_groups_by_rubrique_in_reading_order_and_merges_pages_of_the_same_paper():
    structured = build_structured(
        entries=[
            {"source": "LE FINANCIER D'AFRIQUE", "rubrique": "economie", "headline": "", "summary": "Accord BEAC.", "documents": ["a.jpg"]},
            {"source": "CAMEROON TRIBUNE", "rubrique": "generale", "headline": "« Une volonté de transparence »", "summary": "Le quotidien ouvre sur les élections.", "documents": ["b.jpg"]},
            {"source": "Cameroon Tribune", "rubrique": "generale", "headline": "", "summary": "En page 3, la CAN féminine.", "documents": ["c.jpg"]},
            {"source": "LE JOUR", "rubrique": "generale", "headline": "", "summary": "Paludisme à Yaoundé.", "documents": ["d.jpg"]},
        ],
        country="CAMEROUN",
        date_label="Jeudi 5 mars 2026",
    )

    assert [s["key"] for s in structured["sections"]] == ["generale", "economie"]
    generale = structured["sections"][0]["entries"]
    assert [e["source"] for e in generale] == ["CAMEROON TRIBUNE", "LE JOUR"]
    assert generale[0]["headline"] == "Une volonté de transparence"  # guillemets du modèle retirés
    assert generale[0]["summary"] == "Le quotidien ouvre sur les élections. En page 3, la CAN féminine."
    assert generale[0]["documents"] == ["b.jpg", "c.jpg"]
    assert entry_paragraph(generale[0]).startswith("« Une volonté de transparence » — Le quotidien")
    assert entry_paragraph(generale[1]) == "Paludisme à Yaoundé."

    view = present(structured)
    assert view["sections"][0]["title"] == "LA PRESSE GÉNÉRALE" and view["sections"][0]["color"] == "#1A237E"
    markdown = to_markdown(structured)
    assert markdown.startswith("# REVUE DE PRESSE — CAMEROUN — Jeudi 5 mars 2026")
    assert "### ▶ CAMEROON TRIBUNE" in markdown and "PRESSE ÉCONOMIQUE ET FINANCIÈRE" in markdown


def test_pages_of_a_general_daily_classified_by_topic_stay_together_in_general_press():
    """Constaté en réel : la une politique du Cameroon Tribune classée « politique » et sa page
    santé classée « société » donnaient deux notices dans deux rubriques."""
    structured = build_structured(
        entries=[
            {"source": "CAMEROON TRIBUNE", "rubrique": "politique", "headline": "Elecam et l'ONU", "summary": "Le quotidien ouvre sur les élections.", "documents": ["une.jpg"]},
            {"source": "THE GUARDIAN POST", "rubrique": "generale", "headline": "", "summary": "Décès du PCA de la CDC.", "documents": ["gp.jpg"]},
            {"source": "CAMEROON TRIBUNE", "rubrique": "generale", "headline": "", "summary": "Yaoundé a pris en charge **24 949 cas** de paludisme.", "documents": ["p5.jpg"]},
        ]
    )

    assert [s["key"] for s in structured["sections"]] == ["generale"]
    tribune, guardian = structured["sections"][0]["entries"]
    assert tribune["source"] == "CAMEROON TRIBUNE" and guardian["source"] == "THE GUARDIAN POST"
    assert tribune["summary"] == "Le quotidien ouvre sur les élections. Yaoundé a pris en charge 24 949 cas de paludisme."
    assert tribune["documents"] == ["une.jpg", "p5.jpg"]


def test_reviews_generated_before_the_format_are_not_presented_as_structured():
    assert present({}) is None
    assert present({"sections": []}) is None


# --------------------------------------------------------------------------- pipeline


class _NewspaperClient(FakeMistralClient):
    """Notice par journal selon le nom du fichier ; compte les appels « press_entry »."""

    def __init__(self, broken=()):
        self.entry_calls = []
        self.broken = set(broken)

    def complete_json(self, **kwargs):
        if kwargs.get("response_hint") == "press_entry":
            prompt = kwargs["user_prompt"]
            name = prompt.split("Fichier : ", 1)[1].split("\n", 1)[0]
            self.entry_calls.append(name)
            if name in self.broken:
                return {"resume": ""}  # illisible, y compris à la relance
            papers = {
                "tribune-p1.pdf": ("Cameroon Tribune", "La presse générale", "Une volonté de transparence", "Le quotidien gouvernemental ouvre sur les élections."),
                "tribune-p2.pdf": ("CAMEROON TRIBUNE", "generale", "", "Le journal évoque aussi la CAN féminine."),
                "financier.pdf": ("LE FINANCIER D'AFRIQUE", "économie", "", "Le journal éclaire l'accord BEAC."),
            }
            journal, rubrique, titre, resume = papers[name]
            return {"journal": journal, "rubrique": rubrique, "titre": titre, "resume": resume}
        return super().complete_json(**kwargs)


def _job(organization, user, names):
    job = PressReviewJob.objects.create(organization=organization, created_by=user, title="Revue de presse du 5 mars 2026")
    for index, name in enumerate(names):
        SourceDocument.objects.create(
            organization=organization, job=job, uploaded_by=user, position=index, original_filename=name, kind="pdf",
            file=SimpleUploadedFile(name, make_minimal_pdf_bytes([f"Texte du document {name} assez long pour etre lu."]), content_type="application/pdf"),
        )
    return job


@pytest.mark.django_db
def test_pipeline_writes_one_notice_per_newspaper_grouped_by_rubrique(organization, user, monkeypatch):
    DetectionCategory.objects.create(name="Point important", type="important_point", is_active=True)
    client = _NewspaperClient()
    monkeypatch.setattr(pipeline_tasks, "get_mistral_client", lambda: client)
    job = _job(organization, user, ["tribune-p1.pdf", "financier.pdf", "tribune-p2.pdf"])

    run_press_review_pipeline(str(job.id))

    job.refresh_from_db()
    assert job.status == PressReviewJob.Status.REVIEW_READY, job.error_message
    review = GeneratedReview.objects.get(job=job)
    sections = review.structured["sections"]
    assert [s["key"] for s in sections] == ["generale", "economie"]
    tribune = sections[0]["entries"]
    assert len(tribune) == 1 and tribune[0]["source"] == "CAMEROON TRIBUNE"
    assert tribune[0]["summary"].endswith("Le journal évoque aussi la CAN féminine.")
    assert review.structured["country"] == "CAMEROUN"
    assert review.summary_markdown.startswith("# REVUE DE PRESSE — CAMEROUN — ")
    assert client.entry_calls == ["tribune-p1.pdf", "financier.pdf", "tribune-p2.pdf"]


@pytest.mark.django_db
def test_an_unreadable_notice_keeps_the_document_summaries_instead_of_failing(organization, user, monkeypatch):
    DetectionCategory.objects.create(name="Point important", type="important_point", is_active=True)
    monkeypatch.setattr(pipeline_tasks, "get_mistral_client", lambda: _NewspaperClient(broken={"financier.pdf"}))
    job = _job(organization, user, ["tribune-p1.pdf", "financier.pdf"])

    run_press_review_pipeline(str(job.id))

    job.refresh_from_db()
    assert job.status == PressReviewJob.Status.REVIEW_READY, job.error_message
    entries = [e for s in job.review.structured["sections"] for e in s["entries"]]
    fallback = next(e for e in entries if e["documents"] == ["financier.pdf"])
    assert fallback["source"] == "JOURNAL NON IDENTIFIÉ (document n° 2)"
    assert fallback["summary"]  # résumés de l'étape précédente


@pytest.mark.django_db
def test_a_retried_drafting_step_reuses_the_notices_already_written(organization, user, monkeypatch):
    DetectionCategory.objects.create(name="Point important", type="important_point", is_active=True)
    client = _NewspaperClient()
    monkeypatch.setattr(pipeline_tasks, "get_mistral_client", lambda: client)
    job = _job(organization, user, ["tribune-p1.pdf", "financier.pdf"])
    run_press_review_pipeline(str(job.id))
    first_doc = job.documents.get(original_filename="tribune-p1.pdf")

    # Simule une tentative interrompue : seule la 1re notice était déjà rédigée.
    step = PipelineStepRun.objects.get(job=job, step_name=PipelineStepRun.StepName.REVIEW_DRAFTING)
    step.result_summary = {"entries": {str(first_doc.id): step.result_summary["entries"][str(first_doc.id)]}}
    step.save()
    PressReviewJob.objects.filter(pk=job.pk).update(status=PressReviewJob.Status.HIGHLIGHT_DETECTION)
    client.entry_calls.clear()

    pipeline_tasks.task_draft_review.apply(args=[str(job.id)]).get()

    assert client.entry_calls == ["financier.pdf"]


# --------------------------------------------------------------------------- exports et API


@pytest.fixture
def structured_review(organization, user):
    job = PressReviewJob.objects.create(organization=organization, created_by=user, title="Revue de presse du 5 mars 2026")
    structured = build_structured(
        entries=[
            {"source": "CAMEROON TRIBUNE", "rubrique": "generale", "headline": "Une volonté de transparence",
             "summary": "Le quotidien ouvre sur <script>alert(1)</script> les élections.", "documents": ["a.jpg"]},
            {"source": "LA LUNE", "rubrique": "economie", "headline": "800 FCFA le kg de cacao !",
             "summary": "Le journal interroge.", "documents": ["b.jpg"]},
        ],
        country="CAMEROUN",
        date_label="Jeudi 5 mars 2026",
    )
    return GeneratedReview.objects.create(job=job, structured=structured, summary_markdown=to_markdown(structured))


@pytest.mark.django_db
def test_pdf_follows_the_press_review_layout_and_escapes_the_text(structured_review):
    html = _build_pdf_html(structured_review)

    assert "class='press-masthead'" in html and "CAMEROUN — Jeudi 5 mars 2026" in html
    assert "LA PRESSE GÉNÉRALE" in html and "background:#1A237E" in html
    assert "PRESSE ÉCONOMIQUE ET FINANCIÈRE" in html and "background:#1B5E20" in html
    assert "▶ CAMEROON TRIBUNE" in html
    assert "« Une volonté de transparence » — Le quotidien" in html
    assert "<script>" not in html and "&lt;script&gt;" in html


@pytest.mark.django_db
def test_docx_follows_the_press_review_layout(structured_review, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    _render_docx(structured_review)

    structured_review.refresh_from_db()
    document = Document(io.BytesIO(structured_review.docx_file.read()))
    banners = [t.rows[0].cells[0] for t in document.tables]
    assert banners[0].paragraphs[0].text == "REVUE DE PRESSE"
    assert banners[0].paragraphs[1].text == "CAMEROUN  —  Jeudi 5 mars 2026"
    assert "1A237E" in banners[0]._tc.xml
    assert "LA PRESSE GÉNÉRALE" in banners[1].paragraphs[0].text
    heading = next(p for p in document.paragraphs if p.text == "▶ CAMEROON TRIBUNE")
    assert heading.runs[0].bold and str(heading.runs[0].font.color.rgb) == "8B0000"
    body = next(p for p in document.paragraphs if p.text.startswith("« Une volonté de transparence » — "))
    assert body.alignment == 3  # justifié


@pytest.mark.django_db
def test_review_api_returns_resolved_rubriques(api_client, user, structured_review):
    api_client.credentials(HTTP_AUTHORIZATION=f"Bearer {RefreshToken.for_user(user).access_token}")

    response = api_client.get(f"/api/jobs/{structured_review.job_id}/review/")

    assert response.status_code == 200
    structured = response.data["structured"]
    assert structured["date_label"] == "Jeudi 5 mars 2026"
    assert [s["title"] for s in structured["sections"]] == ["LA PRESSE GÉNÉRALE", "PRESSE ÉCONOMIQUE ET FINANCIÈRE"]
    assert structured["sections"][0]["entries"][0]["source"] == "CAMEROON TRIBUNE"


def test_known_general_dailies_always_land_in_general_press():
    """Le modèle classe parfois une page par son sujet (un décès en une -> « Décès »)."""
    structured = build_structured(
        entries=[{"source": "THE GUARDIAN POST", "rubrique": "deces", "headline": "No hope for Hope!", "summary": "Décès du PCA de la CDC.", "documents": []}]
    )
    assert [s["key"] for s in structured["sections"]] == ["generale"]


def test_a_notice_split_into_several_objects_is_glued_back_together():
    """Constaté en réel : la notice renvoyée en deux objets (sujet principal, puis « En marge »)."""
    from apps.ai_pipeline.schemas import PressEntryResult

    result = PressEntryResult.model_validate([
        {"journal": "MYHOSPITAL", "rubrique": "societe", "titre": "Octobre Rose", "resume": "Le quotidien *Myhospital* met en avant le dépistage."},
        {"journal": "MYHOSPITAL", "rubrique": "societe", "titre": "", "resume": "En marge, des consultations gratuites à Yassa."},
    ])

    assert result.journal == "MYHOSPITAL" and result.titre == "Octobre Rose"
    assert result.resume == "Le quotidien *Myhospital* met en avant le dépistage. En marge, des consultations gratuites à Yassa."
    structured = build_structured(entries=[{"source": "MYHOSPITAL", "rubrique": "societe", "headline": "", "summary": result.resume, "documents": []}])
    assert structured["sections"][0]["entries"][0]["summary"].startswith("Le quotidien Myhospital met")


# --------------------------------------------------------------------------- classement des journaux


def _sources(view, key):
    return [e["source"] for s in view["sections"] if s["key"] == key for e in s["entries"]]


def test_newspapers_are_ranked_from_most_read_to_least_known_within_each_rubric():
    structured = build_structured(
        entries=[
            {"source": "THE SUN", "rubrique": "generale", "headline": "", "summary": "a", "documents": []},
            {"source": "LE PETIT ÉCHO", "rubrique": "generale", "headline": "", "summary": "b", "documents": []},
            {"source": "JOURNAL NON IDENTIFIÉ (photo n° 3)", "rubrique": "generale", "headline": "", "summary": "c", "documents": []},
            {"source": "Mutations", "rubrique": "generale", "headline": "", "summary": "d", "documents": []},
            {"source": "CAMEROON TRIBUNE", "rubrique": "generale", "headline": "", "summary": "e", "documents": []},
            {"source": "LA GAZETTE DU NORD", "rubrique": "generale", "headline": "", "summary": "f", "documents": []},
            {"source": "LE JOUR", "rubrique": "generale", "headline": "", "summary": "g", "documents": []},
        ]
    )

    assert _sources(present(structured), "generale") == [
        "CAMEROON TRIBUNE", "LE JOUR", "MUTATIONS", "THE SUN",
        "LE PETIT ÉCHO", "LA GAZETTE DU NORD",          # inconnus : ordre de lecture
        "JOURNAL NON IDENTIFIÉ (photo n° 3)",          # toujours en dernier
    ]
    assert "## 📰 LA PRESSE GÉNÉRALE\n\n### ▶ CAMEROON TRIBUNE" in to_markdown(structured)


def test_reviews_stored_in_reading_order_are_ranked_when_displayed():
    """Le classement s'applique au rendu : les revues générées avant lui en profitent aussi."""
    stored = {
        "format": "press_review_v1", "country": "CAMEROUN", "date_label": "",
        "sections": [{"key": "generale", "entries": [
            {"source": "THE GUARDIAN POST", "headline": "", "summary": "x", "documents": []},
            {"source": "LE MESSAGER", "headline": "", "summary": "y", "documents": []},
            {"source": "CAMEROON TRIBUNE", "headline": "", "summary": "z", "documents": []},
        ]}],
    }

    assert _sources(present(stored), "generale") == ["CAMEROON TRIBUNE", "LE MESSAGER", "THE GUARDIAN POST"]


def test_a_misread_newspaper_name_is_corrected_merged_and_filed_in_its_rubric():
    """Constaté en réel : « L'OCIL DU SAHEL » (Œ mal lu) ; un titre économique dont la une
    parle d'élections reste en presse économique."""
    structured = build_structured(
        entries=[
            {"source": "L'OCIL DU SAHEL", "rubrique": "generale", "headline": "", "summary": "Attaques.", "documents": ["a.jpg"]},
            {"source": "LE FINANCIER D'AFRIQUE", "rubrique": "politique", "headline": "", "summary": "Accord.", "documents": ["b.jpg"]},
            {"source": "L'ŒIL DU SAHEL", "rubrique": "regionale", "headline": "", "summary": "Bouba Ndjidda.", "documents": ["c.jpg"]},
            {"source": "MUTATIONS (30 ANS)", "rubrique": "societe", "headline": "", "summary": "Visas.", "documents": ["d.jpg"]},
        ]
    )

    view = present(structured)
    assert _sources(view, "regionale") == ["L'ŒIL DU SAHEL"]
    sahel = view["sections"][[s["key"] for s in view["sections"]].index("regionale")]["entries"][0]
    assert sahel["summary"] == "Attaques. Bouba Ndjidda." and sahel["documents"] == ["a.jpg", "c.jpg"]
    assert _sources(view, "economie") == ["LE FINANCIER D'AFRIQUE"]
    # début de nom reconnu : classé et rangé comme Mutations, mais nommé tel qu'imprimé
    assert _sources(view, "generale") == ["MUTATIONS (30 ANS)"]


@pytest.mark.parametrize(
    "printed, expected",
    [
        ("Cameroon Tribune", "CAMEROON TRIBUNE"),
        ("CAMEROUN TRIBUNE", "CAMEROON TRIBUNE"),
        ("LA NOUVELLE EXPRESSI0N", "LA NOUVELLE EXPRESSION"),   # zéro lu à la place du O
        ("THE POST", "THE POST"),                               # pas confondu avec The Guardian Post
        ("LE CITOYEN", "LE CITOYEN"),                           # inconnu : laissé tel quel
    ],
)
def test_newspaper_names_are_normalised_without_confusing_titles(printed, expected):
    from apps.ai_pipeline.media import display_name

    assert display_name(printed) == expected
